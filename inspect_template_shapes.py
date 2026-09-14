from pptx import Presentation

prs = Presentation("SIH2026-IDEA-Presentation-Format.pptx")
print(f"Slide dimensions: {prs.slide_width.inches:.2f} x {prs.slide_height.inches:.2f} inches")
for i, slide in enumerate(prs.slides):
    print(f"\n--- Slide {i+1} ({len(slide.shapes)} shapes) ---")
    for s in slide.shapes:
        text = ""
        if s.has_text_frame:
            text = " | ".join(p.text.strip() for p in s.text_frame.paragraphs if p.text.strip())
        print(f"  [{s.name}] type={s.shape_type} pos=({s.left.inches:.2f}, {s.top.inches:.2f}) size=({s.width.inches:.2f} x {s.height.inches:.2f}) text='{text[:60]}'")
