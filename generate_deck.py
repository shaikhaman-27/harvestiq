import os
import sys
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

INPUT_PPTX = 'SIH2026-IDEA-Presentation-Format.pptx'
OUTPUT_PPTX = 'SIH_26005_Smart_Solar_Cold_Storage_FINAL.pptx'
HERO_IMG = 'hero_cold_storage.jpg'

# Official & Premium Palette
NAVY_HEADER   = RGBColor(15, 23, 42)      # #0F172A
NAVY_TEXT     = RGBColor(30, 41, 59)      # #1E293B
BLUE_SIH      = RGBColor(0, 112, 192)     # #0070C0 Official SIH Blue
BLUE_DARK     = RGBColor(3, 105, 161)     # #0369A1
BLUE_LIGHT    = RGBColor(239, 246, 255)   # #EFF6FF
BLUE_BORDER   = RGBColor(186, 230, 253)   # #BAE6FD
GREEN_DARK    = RGBColor(22, 101, 52)     # #166534
GREEN_LIGHT   = RGBColor(240, 253, 244)   # #F0FDF4
GREEN_BORDER  = RGBColor(134, 239, 172)   # #86EFAC
AMBER_DARK    = RGBColor(180, 83, 9)      # #B45309
AMBER_LIGHT   = RGBColor(254, 243, 199)   # #FEF3C7
AMBER_BORDER  = RGBColor(252, 211, 77)    # #FCD34D
RED_DARK      = RGBColor(185, 28, 28)     # #B91C1C
RED_LIGHT     = RGBColor(254, 242, 242)   # #FEF2F2
RED_BORDER    = RGBColor(252, 165, 165)   # #FCA5A5
PURPLE_DARK   = RGBColor(109, 40, 217)    # #6D28D9
PURPLE_LIGHT  = RGBColor(245, 243, 255)   # #F5F3FF
PURPLE_BORDER = RGBColor(221, 214, 254)   # #DDD6FE

WHITE         = RGBColor(255, 255, 255)
CARD_BG_MUTED = RGBColor(248, 250, 252)   # #F8FAFC
BORDER_MUTED  = RGBColor(203, 213, 225)   # #CBD5E1
BORDER_LIGHT  = RGBColor(226, 232, 240)   # #E2E8F0
TEXT_DARK     = RGBColor(15, 23, 42)      # #0F172A
TEXT_BODY     = RGBColor(51, 65, 85)      # #334155
TEXT_MUTED    = RGBColor(100, 116, 139)   # #64748B

FONT_HEAD = 'Segoe UI'
FONT_BODY = 'Arial'

def add_card(slide, left, top, width, height, bg_color=WHITE, border_color=BORDER_MUTED, border_width=1, shape_type=MSO_SHAPE.ROUNDED_RECTANGLE):
    card = slide.shapes.add_shape(shape_type, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    if border_color:
        card.line.color.rgb = border_color
        card.line.width = Pt(border_width)
    else:
        card.line.fill.background()
    return card

def add_rect(slide, left, top, width, height, bg_color=WHITE, border_color=BORDER_MUTED, border_width=1):
    return add_card(slide, left, top, width, height, bg_color, border_color, border_width, MSO_SHAPE.RECTANGLE)

def add_badge(slide, left, top, width, height, text, bg_color=BLUE_SIH, text_color=WHITE, font_size=8.5, bold=True, shape_type=MSO_SHAPE.ROUNDED_RECTANGLE):
    badge = slide.shapes.add_shape(shape_type, left, top, width, height)
    badge.fill.solid()
    badge.fill.fore_color.rgb = bg_color
    badge.line.fill.background()
    tf = badge.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = Inches(0.02)
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = FONT_HEAD
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = text_color
    p.alignment = PP_ALIGN.CENTER
    return badge

def add_textbox(slide, left, top, width, height, margin=0.03):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = Inches(margin)
    return tb

def setup_header(slide, title_text, subtitle_text):
    # Setup Title 1 placeholder
    for s in slide.shapes:
        if s.name == 'Title 1' and s.has_text_frame:
            s.left = Inches(1.80)
            s.top = Inches(0.08)
            s.width = Inches(8.80)
            s.height = Inches(0.48)
            tf = s.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = title_text
            p.font.name = FONT_HEAD
            p.font.size = Pt(20)
            p.font.bold = True
            p.font.color.rgb = NAVY_HEADER
            p.alignment = PP_ALIGN.LEFT
    
    # Subtitle directly under title
    tb_sub = add_textbox(slide, Inches(1.80), Inches(0.56), Inches(8.80), Inches(0.32), margin=0.01)
    p2 = tb_sub.text_frame.paragraphs[0]
    p2.text = subtitle_text
    p2.font.name = FONT_HEAD
    p2.font.size = Pt(9.5)
    p2.font.bold = True
    p2.font.color.rgb = BLUE_DARK

def format_team_oval_and_footers(slide):
    for s in slide.shapes:
        if 'Oval' in s.name and s.has_text_frame:
            s.left = Inches(0.35)
            s.top = Inches(0.12)
            s.width = Inches(1.30)
            s.height = Inches(0.65)
            s.fill.solid()
            s.fill.fore_color.rgb = BLUE_SIH
            s.line.fill.background()
            for p in s.text_frame.paragraphs:
                p.text = 'HarvestIQ'
                p.font.name = FONT_HEAD
                p.font.size = Pt(11)
                p.font.bold = True
                p.font.color.rgb = WHITE
                p.alignment = PP_ALIGN.CENTER
        if 'Footer Placeholder' in s.name and s.has_text_frame:
            for p in s.text_frame.paragraphs:
                p.font.name = FONT_BODY
                p.font.size = Pt(9)
                p.font.color.rgb = WHITE
        if 'Slide Number' in s.name and s.has_text_frame:
            for p in s.text_frame.paragraphs:
                p.font.name = FONT_BODY
                p.font.size = Pt(10)
                p.font.bold = True
                p.font.color.rgb = WHITE

def clear_old_placeholders(slide):
    for s in list(slide.shapes):
        if s.has_text_frame:
            txt = s.text_frame.text.lower()
            if 'proposed solution (describe' in txt or 'technologies to be used' in txt or \
               'analysis of the feasibility' in txt or 'potential impact on the target' in txt or \
               'details / links of the reference' in txt:
                sp = s._element
                sp.getparent().remove(sp)

print('generate_deck.py Part 1 ready')
