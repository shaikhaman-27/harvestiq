from pptx import Presentation

prs = Presentation("SIH2026-IDEA-Presentation-Format.pptx")
print("Original slides:", len(prs.slides))

# delete 7th slide (index 6)
rId = prs.slides._sldIdLst[6].rId
prs.part.drop_rel(rId)
del prs.slides._sldIdLst[6]

prs.save("test_6slides.pptx")

# reload and verify
prs2 = Presentation("test_6slides.pptx")
print("New slides count:", len(prs2.slides))
for i, s in enumerate(prs2.slides):
    print(f"Slide {i+1}")
