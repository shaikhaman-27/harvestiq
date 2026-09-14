import pptx

prs = pptx.Presentation('SIH2026-IDEA-Presentation-Format.pptx')
with open('template_details.txt', 'w', encoding='utf-8') as f:
    for i, slide in enumerate(prs.slides):
        f.write(f"=== SLIDE {i+1} ===\n")
        for s in slide.shapes:
            f.write(f"  Shape: {s.name} (type: {s.shape_type})\n")
            f.write(f"    Left: {s.left}, Top: {s.top}, Width: {s.width}, Height: {s.height}\n")
            if s.has_text_frame:
                for p in s.text_frame.paragraphs:
                    for r in p.runs:
                        try:
                            c = r.font.color.rgb if r.font.color and r.font.color.type == 1 else "theme/auto"
                        except Exception:
                            c = "unknown"
                        f.write(f"      Text: '{r.text}' | Font: {r.font.name} | Size: {r.font.size} | Color: {c}\n")

print("Template details written to template_details.txt")
