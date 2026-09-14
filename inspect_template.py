import pptx

prs = pptx.Presentation('SIH2026-IDEA-Presentation-Format.pptx')
with open('inspect_output.txt', 'w', encoding='utf-8') as out:
    for i, slide in enumerate(prs.slides):
        out.write(f'=== SLIDE {i+1} (shapes: {len(slide.shapes)}) ===\n')
        for s_idx, s in enumerate(slide.shapes):
            out.write(f'  Shape {s_idx}: name="{s.name}", type={s.shape_type}, pos=({s.left},{s.top}), size=({s.width},{s.height})\n')
            if s.has_text_frame:
                tf = s.text_frame
                for p_idx, p in enumerate(tf.paragraphs):
                    txt = p.text.strip()
                    if txt:
                        out.write(f'     Para {p_idx} (level={p.level}): "{txt}"\n')
print('Inspection written to inspect_output.txt')
