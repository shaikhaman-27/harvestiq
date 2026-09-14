import os
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

INPUT_PPTX = "SIH2026-IDEA-Presentation-Format.pptx"
OUTPUT_PPTX = "SIH_26005_Smart_Solar_Cold_Storage_FINAL.pptx"
HERO_IMG = "hero_cold_storage.jpg"

# Color Palette
NAVY_TITLE   = RGBColor(15, 23, 42)      # #0F172A Dark Charcoal / Deep Slate
NAVY_CARD    = RGBColor(30, 41, 59)      # #1E293B Deep Navy Card
BLUE_SIH     = RGBColor(0, 112, 192)     # #0070C0 Official SIH Blue
BLUE_DARK    = RGBColor(3, 105, 161)     # #0369A1 Dark Blue Accent
BLUE_LIGHT   = RGBColor(238, 246, 255)   # #EEF6FF Soft Blue Tint
GREEN_AGRI   = RGBColor(22, 101, 52)     # #166534 Agritech Forest Green
GREEN_EMERALD= RGBColor(16, 149, 106)    # #10956A Vibrant Green
GREEN_TINT   = RGBColor(240, 253, 244)   # #F0FDF4 Soft Green Card BG
AMBER_SOLAR  = RGBColor(217, 119, 6)    # #D97706 Solar Amber
AMBER_TINT   = RGBColor(254, 243, 199)   # #FEF3C7 Soft Amber BG
TEAL_COOL    = RGBColor(13, 148, 136)    # #0D9488 Cool Teal
TEAL_TINT    = RGBColor(240, 253, 250)   # #F0FDFA Soft Teal BG
CARD_BG      = RGBColor(248, 250, 252)   # #F8FAFC Crisp Slate Card BG
CARD_BORDER  = RGBColor(203, 213, 225)   # #CBD5E1 Slate Border
WHITE        = RGBColor(255, 255, 255)
TEXT_DARK    = RGBColor(30, 41, 59)      # #1E293B Body Text
TEXT_MUTED   = RGBColor(100, 116, 139)   # #64748B Secondary Text
RED_ACCENT   = RGBColor(185, 28, 28)     # #B91C1C Problem/Risk Red
RED_TINT     = RGBColor(254, 242, 242)   # #FEF2F2 Soft Red BG

FONT_HEAD = "Segoe UI"
FONT_BODY = "Arial"

def clear_instruction_boxes(slide):
    """Removes the template instruction text boxes so we can place custom rich cards."""
    for s in list(slide.shapes):
        if s.has_text_frame:
            txt = s.text_frame.text.lower()
            if "proposed solution (describe" in txt or "technologies to be used" in txt or \
               "analysis of the feasibility" in txt or "potential impact on the target" in txt or \
               "details / links of the reference" in txt:
                sp = s._element
                sp.getparent().remove(sp)

def remove_shape_by_name(slide, shape_name):
    for s in list(slide.shapes):
        if s.name == shape_name:
            sp = s._element
            sp.getparent().remove(sp)

def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER, border_width=1):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    if border_color:
        card.line.color.rgb = border_color
        card.line.width = Pt(border_width)
    else:
        card.line.fill.background()
    return card

def add_rect(slide, left, top, width, height, bg_color=CARD_BG, border_color=None, border_width=1):
    rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    rect.fill.solid()
    rect.fill.fore_color.rgb = bg_color
    if border_color:
        rect.line.color.rgb = border_color
        rect.line.width = Pt(border_width)
    else:
        rect.line.fill.background()
    return rect

def add_badge(slide, left, top, width, height, text, bg_color=BLUE_SIH, text_color=WHITE, font_size=8.5, bold=True):
    badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    badge.fill.solid()
    badge.fill.fore_color.rgb = bg_color
    badge.line.fill.background()
    tf = badge.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = Inches(0.02)
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = FONT_HEAD
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = text_color
    p.alignment = PP_ALIGN.CENTER
    return badge

def add_tb(slide, left, top, width, height, text="", font_name=FONT_BODY, font_size=10, bold=False, color=TEXT_DARK, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = Inches(0.04)
    if text:
        p = tf.paragraphs[0]
        p.text = text
        p.font.name = font_name
        p.font.size = Pt(font_size)
        p.font.bold = bold
        p.font.color.rgb = color
        p.alignment = align
    return tb

def add_header(slide, main_title, subtitle, section_badge):
    # Section Badge
    add_badge(slide, Inches(1.85), Inches(0.18), Inches(2.7), Inches(0.28), section_badge, BLUE_SIH, WHITE, font_size=8.5, bold=True)
    
    # Main Header
    tb = add_tb(slide, Inches(4.65), Inches(0.12), Inches(6.0), Inches(0.42))
    p = tb.text_frame.paragraphs[0]
    p.text = main_title
    p.font.name = FONT_HEAD
    p.font.size = Pt(15.5)
    p.font.bold = True
    p.font.color.rgb = NAVY_TITLE
    
    # Subtitle
    tb2 = add_tb(slide, Inches(1.85), Inches(0.52), Inches(8.8), Inches(0.35))
    p2 = tb2.text_frame.paragraphs[0]
    p2.text = subtitle
    p2.font.name = FONT_BODY
    p2.font.size = Pt(10)
    p2.font.color.rgb = TEXT_MUTED

def style_oval_and_footers(slide):
    for s in slide.shapes:
        if "Oval" in s.name and s.has_text_frame:
            s.fill.solid()
            s.fill.fore_color.rgb = RGBColor(241, 245, 249)
            s.line.color.rgb = BLUE_SIH
            s.line.width = Pt(1)
            for p in s.text_frame.paragraphs:
                p.text = "Team ID / Name\n[Registered]"
                p.font.name = FONT_HEAD
                p.font.size = Pt(8)
                p.font.bold = True
                p.font.color.rgb = BLUE_SIH
                p.alignment = PP_ALIGN.CENTER
        if "Footer Placeholder" in s.name and s.has_text_frame:
            for p in s.text_frame.paragraphs:
                p.font.name = FONT_BODY
                p.font.size = Pt(9)
                p.font.color.rgb = WHITE
        if "Slide Number" in s.name and s.has_text_frame:
            for p in s.text_frame.paragraphs:
                p.font.name = FONT_BODY
                p.font.size = Pt(10)
                p.font.bold = True
                p.font.color.rgb = WHITE
        # Remove empty or duplicate titles if needed
        if "Title 1" in s.name and s.has_text_frame:
            # We override with custom header
            s.text_frame.text = ""

print("Helper definitions complete.")
