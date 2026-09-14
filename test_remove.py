from pptx import Presentation

prs = Presentation("SIH2026-IDEA-Presentation-Format.pptx")
slide = prs.slides[0]
for shape in list(slide.shapes):
    if shape.has_text_frame and "Problem Statement ID" in shape.text_frame.text:
        sp = shape._element
        sp.getparent().remove(sp)
        print("Removed placeholder successfully")
