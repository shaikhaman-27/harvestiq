import os
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Paths
INPUT_PPTX = "SIH2026-IDEA-Presentation-Format.pptx"
OUTPUT_PPTX = "SIH_26005_Smart_Solar_Cold_Storage_FINAL.pptx"
HERO_IMG = "hero_cold_storage.jpg"

# Color Palette
NAVY_TITLE = RGBColor(15, 23, 42)      # #0F172A Dark Charcoal/Navy
BLUE_SIH = RGBColor(0, 112, 192)      # #0070C0 SIH Official Blue
BLUE_ACCENT = RGBColor(2, 132, 199)   # #0284C7 Sky Blue
GREEN_AGRI = RGBColor(22, 101, 52)    # #166534 Deep Green
GREEN_MINT = RGBColor(240, 253, 244)  # #F0FDF4 Soft Green BG
AMBER_SOLAR = RGBColor(217, 119, 6)   # #D97706 Solar Amber
TEAL_COOL = RGBColor(13, 148, 136)    # #0D9488 Cool Teal
CARD_BG = RGBColor(248, 250, 252)     # #F8FAFC Clean Card BG
CARD_BORDER = RGBColor(203, 213, 225) # #CBD5E1 Slate Border
WHITE = RGBColor(255, 255, 255)
TEXT_DARK = RGBColor(30, 41, 59)      # #1E293B Body Text
TEXT_MUTED = RGBColor(100, 116, 139)  # #64748B Subtitle/Muted
DARK_BG = RGBColor(15, 23, 42)

FONT_HEADING = "Segoe UI"
FONT_BODY = "Arial"

def clear_shape_text(shape):
    if shape.has_text_frame:
        shape.text_frame.text = ""

def remove_shape_by_name(slide, shape_name):
    for s in list(slide.shapes):
        if s.name == shape_name:
            sp = s._element
            sp.getparent().remove(sp)

def add_rounded_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER, border_width=1):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    if border_color:
        card.line.color.rgb = border_color
        card.line.width = Pt(border_width)
    else:
        card.line.fill.background()
    return card

def add_rect(slide, left, top, width, height, bg_color=CARD_BG, border_color=None, border_width=1):
    rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    rect.fill.solid()
    rect.fill.fore_color.rgb = bg_color
    if border_color:
        rect.line.color.rgb = border_color
        rect.line.width = Pt(border_width)
    else:
        rect.line.fill.background()
    return rect

def add_badge(slide, left, top, width, height, text, bg_color=BLUE_SIH, text_color=WHITE, font_size=9, bold=True):
    badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    badge.fill.solid()
    badge.fill.fore_color.rgb = bg_color
    badge.line.fill.background()
    tf = badge.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = Inches(0.04)
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = FONT_HEADING
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = text_color
    p.alignment = PP_ALIGN.CENTER
    return badge

def add_textbox(slide, left, top, width, height, text="", font_name=FONT_BODY, font_size=11, bold=False, color=TEXT_DARK, align=PP_ALIGN.LEFT):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = Inches(0.05)
    if text:
        p = tf.paragraphs[0]
        p.text = text
        p.font.name = font_name
        p.font.size = Pt(font_size)
        p.font.bold = bold
        p.font.color.rgb = color
        p.alignment = align
    return txBox

def add_header(slide, main_title, subtitle, section_badge):
    # Header zone: left=0.6", top=0.2", width=9.8"
    add_badge(slide, Inches(0.6), Inches(0.2), Inches(2.5), Inches(0.3), section_badge, BLUE_SIH, WHITE, font_size=9, bold=True)
    
    tb = add_textbox(slide, Inches(3.2), Inches(0.15), Inches(6.8), Inches(0.45))
    p = tb.text_frame.paragraphs[0]
    p.text = main_title
    p.font.name = FONT_HEADING
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = NAVY_TITLE
    
    tb2 = add_textbox(slide, Inches(0.6), Inches(0.55), Inches(9.8), Inches(0.35))
    p2 = tb2.text_frame.paragraphs[0]
    p2.text = subtitle
    p2.font.name = FONT_BODY
    p2.font.size = Pt(11)
    p2.font.color.rgb = TEXT_MUTED

def style_oval_and_footer(slide):
    for s in slide.shapes:
        if "Oval" in s.name and s.has_text_frame:
            # Style team name oval
            s.fill.solid()
            s.fill.fore_color.rgb = RGBColor(241, 245, 249)
            s.line.color.rgb = BLUE_SIH
            s.line.width = Pt(1)
            for p in s.text_frame.paragraphs:
                p.font.name = FONT_HEADING
                p.font.size = Pt(9)
                p.font.bold = True
                p.font.color.rgb = BLUE_SIH
                p.alignment = PP_ALIGN.CENTER
        if "Footer Placeholder" in s.name and s.has_text_frame:
            for p in s.text_frame.paragraphs:
                p.font.name = FONT_BODY
                p.font.size = Pt(9)
                p.font.color.rgb = WHITE
        if "Slide Number" in s.name and s.has_text_frame:
            for p in s.text_frame.paragraphs:
                p.font.name = FONT_BODY
                p.font.size = Pt(10)
                p.font.bold = True
                p.font.color.rgb = WHITE

print("Loaded module successfully")
