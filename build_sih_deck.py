import os
import sys
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Paths
INPUT_PPTX = "SIH2026-IDEA-Presentation-Format.pptx"
OUTPUT_PPTX = "SIH_26005_Smart_Solar_Cold_Storage_FINAL.pptx"
HERO_IMG = "hero_cold_storage.jpg"

# Color Palette
NAVY_TITLE    = RGBColor(15, 23, 42)      # #0F172A
NAVY_CARD     = RGBColor(30, 41, 59)      # #1E293B
BLUE_SIH      = RGBColor(0, 112, 192)     # #0070C0 Official SIH Blue
BLUE_DARK     = RGBColor(2, 132, 199)     # #0284C7
BLUE_LIGHT    = RGBColor(239, 246, 255)   # #EFF6FF
BLUE_BORDER   = RGBColor(186, 230, 253)   # #BAE6FD
GREEN_DARK    = RGBColor(22, 101, 52)     # #166534
GREEN_LIGHT   = RGBColor(240, 253, 244)   # #F0FDF4
GREEN_BORDER  = RGBColor(134, 239, 172)   # #86EFAC
AMBER_DARK    = RGBColor(180, 83, 9)      # #B45309
AMBER_LIGHT   = RGBColor(254, 243, 199)   # #FEF3C7
AMBER_BORDER  = RGBColor(252, 211, 77)    # #FCD34D
RED_DARK      = RGBColor(185, 28, 28)     # #B91C1C
RED_LIGHT     = RGBColor(254, 242, 242)   # #FEF2F2
RED_BORDER    = RGBColor(252, 165, 165)   # #FCA5A5
CARD_BG       = RGBColor(255, 255, 255)   # White
CARD_BG_MUTED = RGBColor(248, 250, 252)   # #F8FAFC
BORDER_MUTED  = RGBColor(203, 213, 225)   # #CBD5E1
BORDER_LIGHT  = RGBColor(226, 232, 240)   # #E2E8F0
WHITE         = RGBColor(255, 255, 255)
TEXT_DARK     = RGBColor(15, 23, 42)      # #0F172A
TEXT_BODY     = RGBColor(51, 65, 85)      # #334155
TEXT_MUTED    = RGBColor(100, 116, 139)   # #64748B

FONT_HEAD = "Segoe UI"
FONT_BODY = "Arial"

def add_card(slide, left, top, width, height, bg_color=WHITE, border_color=BORDER_MUTED, border_width=1):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    if border_color:
        card.line.color.rgb = border_color
        card.line.width = Pt(border_width)
    else:
        card.line.fill.background()
    return card

def add_rect(slide, left, top, width, height, bg_color=WHITE, border_color=BORDER_MUTED, border_width=1):
    card = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    if border_color:
        card.line.color.rgb = border_color
        card.line.width = Pt(border_width)
    else:
        card.line.fill.background()
    return card

def add_badge(slide, left, top, width, height, text, bg_color=BLUE_SIH, text_color=WHITE, font_size=9, bold=True):
    badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    badge.fill.solid()
    badge.fill.fore_color.rgb = bg_color
    badge.line.fill.background()
    tf = badge.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = Inches(0.02)
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = FONT_HEAD
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = text_color
    p.alignment = PP_ALIGN.CENTER
    return badge

def add_textbox(slide, left, top, width, height):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = Inches(0.03)
    return tb

def style_para(p, text, font_name=FONT_BODY, font_size=9, bold=False, color=TEXT_BODY, align=PP_ALIGN.LEFT, space_after=0):
    p.text = text
    p.font.name = font_name
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = align
    p.space_after = Pt(space_after)

def add_run(p, text, font_name=FONT_BODY, font_size=9, bold=False, color=TEXT_BODY):
    run = p.add_run()
    run.text = text
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color
    return run

print("Base helper functions defined")
