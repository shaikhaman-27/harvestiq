import pptx

prs = pptx.Presentation('SIH2026-IDEA-Presentation-Format.pptx')
for i, slide in enumerate(prs.slides):
    print(f"=== SLIDE {i+1} ===")
    for s in slide.shapes:
        info = []
        if hasattr(s, 'fill'):
            try:
                if s.fill.type == 1: # solid
                    info.append(f"fill={s.fill.fore_color.rgb}")
                else:
                    info.append(f"fill_type={s.fill.type}")
            except Exception as e:
                info.append(f"fill_err={e}")
        if hasattr(s, 'line'):
            try:
                if s.line.fill.type == 1:
                    info.append(f"line={s.line.color.rgb}")
            except Exception:
                pass
        print(f"  {s.name} ({s.shape_type}): {', '.join(info)}")
    if i >= 2:
        break
